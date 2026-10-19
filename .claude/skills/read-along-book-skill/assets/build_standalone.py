#!/usr/bin/env python3
"""Bundle the built book into ONE file that works anywhere, with no folder beside it.

Run it from the book folder after build_book.py:
    python3 build_standalone.py                      (file named after the Settings sheet: file_name, else title_en)
    python3 build_standalone.py "My book.html"       (or name it)
Every picture, recording and font is embedded. Fonts come from fonts/; if that folder is missing the
file keeps loading them from Google Fonts (then it needs internet for the fonts only).
"""
import base64, glob, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'index.html')
FONTS = os.path.join(HERE, 'fonts')
MIME = {'.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.png': 'image/png', '.webp': 'image/webp',
        '.mp3': 'audio/mpeg', '.m4a': 'audio/mp4', '.wav': 'audio/wav', '.woff2': 'font/woff2', '.woff': 'font/woff'}


def data_uri(path):
    ext = os.path.splitext(path)[1].lower()
    with open(path, 'rb') as f:
        return f'data:{MIME.get(ext, "application/octet-stream")};base64,' + base64.b64encode(f.read()).decode()


def out_name():
    if len(sys.argv) > 1: return sys.argv[1]
    try:
        from openpyxl import load_workbook
        books = [p for p in glob.glob(os.path.join(HERE, '*.xlsx')) if not os.path.basename(p).startswith('~$')]
        if len(books) == 1:
            s = {}
            for row in load_workbook(books[0], data_only=True)['Settings'].iter_rows(min_row=2, values_only=True):
                if row and row[0]: s[str(row[0]).strip()] = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
            name = s.get('file_name') or s.get('title_en')
            if name: return name if name.lower().endswith('.html') else re.sub(r'[^\w\s-]', '', name).strip() + '.html'
    except Exception:
        pass
    return 'Read-along book.html'


def main():
    if not os.path.exists(SRC): sys.exit('index.html not found; run build_book.py first')
    html = open(SRC, encoding='utf-8').read()
    missing = []
    def swap(m):
        rel = m.group(1); path = os.path.join(HERE, rel)
        if not os.path.exists(path): missing.append(rel); return m.group(0)
        return '"' + data_uri(path) + '"'
    html, n = re.subn(r'"((?:images|audio)/[^"]+)"', swap, html)          # pictures and recordings named in the page data

    css_path = os.path.join(FONTS, 'fonts.css')
    links = re.compile(r'<link rel="preconnect"[^>]*>\s*|<link rel="stylesheet" href="https://fonts\.googleapis\.com[^"]*">\s*')
    if os.path.exists(css_path):
        css = open(css_path, encoding='utf-8').read()
        css = re.sub(r'url\(([^)]+)\)', lambda m: 'url(' + data_uri(os.path.join(FONTS, m.group(1).strip('\'"'))) + ')', css)
        html = links.sub('', html, count=2)                                 # drop the two preconnects
        html = links.sub('<style>\n' + css + '</style>\n', html, count=1)    # the stylesheet link becomes inline @font-face rules
        fonts_note = 'fonts embedded'
    else:
        fonts_note = 'fonts/ folder missing, fonts still load from Google'
    if not html.lstrip().lower().startswith('<!doctype'): html = '<!doctype html>\n' + html
    out = os.path.join(HERE, out_name())
    open(out, 'w', encoding='utf-8').write(html)
    print(f'wrote {os.path.basename(out)}: {n} files embedded, {fonts_note}, {os.path.getsize(out) / 1e6:.1f} MB')
    for m in missing: print(f'warning: file not found, left as a link: {m}')


if __name__ == '__main__':
    main()
