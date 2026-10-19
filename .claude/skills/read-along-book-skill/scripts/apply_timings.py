#!/usr/bin/env python3
"""Move tapped timings into the spreadsheet (the spreadsheet is the source; timings.csv is only a hand-off).

    python3 scripts/apply_timings.py                      timings.csv in the book folder, saved by the page's Time it mode
    python3 scripts/apply_timings.py taps.csv --raw       raw taps (Line,Start[,End]) typed or pasted: each Start is
                                                          moved back to the speech onset and widened by the margin
    cat taps.csv | python3 scripts/apply_timings.py - --raw

Lines not in the file keep their times. The line just before the first updated line gets a new End so the two
meet in the pause. Timing notes in Check notes are cleared for updated lines. timings.csv is then deleted
(so build_book.py uses the spreadsheet).
"""
import argparse, csv, io, os, re, sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import MARGIN, boundary, end_after, find_xlsx, key_of, onset, pause_at, pause_lists

TIMING_NOTE = re.compile(r'\s*(Timing guessed[^.]*\.|Timing: [^.]*\.|Page turn in the movie[^.]*\.(\s*start time may need a nudge\.)?|please check\.|Please listen to the last page\.)', re.I)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('csv', nargs='?', default='timings.csv')
    ap.add_argument('--raw', action='store_true', help='Start values are raw taps, not already snapped by the page')
    ap.add_argument('--xlsx')
    ap.add_argument('--keep-csv', action='store_true')
    a = ap.parse_args()
    folder = os.getcwd()
    xlsx = find_xlsx(folder, a.xlsx)
    text = sys.stdin.read() if a.csv == '-' else open(a.csv, encoding='utf-8-sig').read()
    rows = {}
    for r in csv.reader(io.StringIO(text.strip())):
        if not r or not r[0].strip() or r[0].strip().lower() in ('line', 'timibgs:'): continue
        try: rows[key_of(r[0])] = [float(r[1]), float(r[2]) if len(r) > 2 and r[2].strip() else None]
        except ValueError: continue
    if not rows: sys.exit('no Line,Start,End rows found')

    lists = pause_lists(folder)
    wb = load_workbook(xlsx)
    ws = wb['Sentences']
    order, cells = [], {}
    for row in ws.iter_rows(min_row=2):
        k = key_of(row[0].value)
        if k: order.append(k); cells[k] = row
    todo = [k for k in order if k in rows]
    if not todo: sys.exit(f'none of the lines {list(rows)} are in the spreadsheet')

    if a.raw:                                              # snap each tap to its speech onset, then widen
        starts, prev_end = {}, {}
        for k in todo:
            on = onset(rows[k][0], lists['fine'])
            e, s = boundary(on, lists)
            starts[k] = s; prev_end[k] = e
        for i, k in enumerate(todo):
            nxt = todo[i + 1] if i + 1 < len(todo) else None
            if nxt and order.index(nxt) == order.index(k) + 1: end = prev_end[nxt]
            elif rows[k][1] and rows[k][1] > starts[k] + 0.3: end = rows[k][1]
            else: end = end_after(starts[k], lists)
            rows[k] = [starts[k], end]
    for k in todo:
        if rows[k][1] is None: rows[k][1] = end_after(rows[k][0], lists)

    yellow = PatternFill('solid', fgColor='FFF4B8')
    for k in todo:
        row = cells[k]
        row[2].value, row[3].value = round(rows[k][0], 2), round(rows[k][1], 2)
        note = TIMING_NOTE.sub('', (row[7].value or '')).strip() if len(row) > 7 else ''
        if len(row) > 7: row[7].value = note or None; row[7].fill = yellow if note else PatternFill(fill_type=None)
    # the line before the first updated one should end where the new first line's pause begins
    first = todo[0]; i = order.index(first)
    if i > 0 and order[i - 1] not in todo:
        before = cells[order[i - 1]]
        for name in ('strict', 'coarse', 'fine'):
            p = pause_at(rows[first][0] + MARGIN + 0.01, lists[name])
            if p: before[3].value = round(p[0] + min(MARGIN, max(0.03, (p[1] - p[0]) / 2 - 0.03)), 2); break
    wb.save(xlsx)
    print(f'wrote {len(todo)} lines into {os.path.basename(xlsx)}: ' + ', '.join(f'{k} {rows[k][0]:.2f}-{rows[k][1]:.2f}' for k in todo))
    if a.csv != '-' and os.path.exists(a.csv) and not a.keep_csv and os.path.basename(a.csv) == 'timings.csv':
        os.remove(a.csv); print('removed timings.csv (the spreadsheet now has the times); run python3 build_book.py')


if __name__ == '__main__':
    main()
