#!/usr/bin/env python3
"""First-draft Start/End for every line, from the pauses in the narration, so the book plays before anyone taps along.

    python3 scripts/guess_timings.py                         (all lines; shares the speech out by sentence length)
    python3 scripts/guess_timings.py --turns "5:10,6:35.5"   (page:seconds where a video of the book turns each page;
                                                              page starts snap to the pause nearest each turn)
    python3 scripts/guess_timings.py --only-empty            (leave lines that already have times)

Run scripts/detect_pauses.py first. The guesses are marked "Timing guessed" in Check notes; Time it mode in the page
(then scripts/apply_timings.py) replaces them with real ones.
"""
import argparse, os, re, sys
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import MARGIN, ffprobe_duration, find_xlsx, key_of, pause_lists, syllables


def split_group(segs, weights):
    """Share speech segments [(a,b)] among lines by weight → [(start, end)] in real time, cuts snapped to pauses."""
    T = sum(b - a for a, b in segs)
    if not segs or T <= 0 or not weights: return [(None, None)] * len(weights)
    W = sum(weights)
    def real(pos):                       # speech-axis position → real time
        for a, b in segs:
            if pos <= b - a + 1e-9: return a + max(0, pos)
            pos -= b - a
        return segs[-1][1]
    out, pos = [], 0.0
    gaps = [(segs[i][1], segs[i + 1][0]) for i in range(len(segs) - 1)]          # pauses inside the group
    cut_prev = segs[0][0]
    for i, w in enumerate(weights):
        pos += T * w / W
        if i == len(weights) - 1:
            out.append((cut_prev, max(segs[-1][1], cut_prev + 0.5))); break
        want = max(real(pos), cut_prev + 0.5)                                  # never run backwards
        avg = T / len(weights)
        near = min(gaps, key=lambda g: abs(g[0] - want), default=None)
        if near and abs(near[0] - want) <= max(0.6, 0.45 * avg) and near[0] > cut_prev + 0.3:
            gaps = [g for g in gaps if g[0] > near[0]]
            m = min(MARGIN, max(0.03, (near[1] - near[0]) / 2 - 0.03))
            out.append((cut_prev, near[0] + m)); cut_prev = near[1] - m
        else:
            out.append((cut_prev, want)); cut_prev = want
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--turns', default='', help='page:seconds pairs, comma separated')
    ap.add_argument('--only-empty', action='store_true')
    ap.add_argument('--min-pause', type=float, default=0.45, help='a gap at least this long (s) can separate two lines')
    ap.add_argument('--xlsx')
    a = ap.parse_args()
    folder = os.getcwd()
    xlsx = find_xlsx(folder, a.xlsx)
    lists = pause_lists(folder)
    if not lists['fine']: sys.exit('no pause lists in text/; run scripts/detect_pauses.py first')
    wb = load_workbook(xlsx)
    audio = next((str(r[1].value) for r in wb['Settings'].iter_rows(min_row=2) if r[0].value == 'audio'), 'audio/narration.mp3')
    dur = ffprobe_duration(os.path.join(folder, audio))
    seps = [p for p in (lists['coarse'] or lists['fine']) if p[1] - p[0] >= a.min_pause]
    segs, t = [], 0.0
    for p in seps:
        if p[0] > t + 0.2: segs.append((t, p[0]))
        t = p[1]
    if dur > t + 0.3: segs.append((t, dur))

    ws = wb['Sentences']
    rows = [(key_of(r[0].value), r) for r in ws.iter_rows(min_row=2) if key_of(r[0].value)]
    title = next((r for k, r in rows if k == 'title'), None)
    lines = [(k, r) for k, r in rows if k != 'title']
    if title is not None and segs and segs[0][1] < 15 and (len(segs) == 1 or segs[1][0] - segs[0][1] >= 0.8):
        if not (a.only_empty and title[3].value): title[2].value, title[3].value = 0, round(segs[0][1] + MARGIN, 2)
        segs = segs[1:]

    # groups: whole book, or one group per page when turns are known
    turns = {}
    for tok in a.turns.split(','):
        if ':' in tok: p, s = tok.split(':'); turns[key_of(p)] = float(s)
    groups = []                                            # (lines, segs)
    if turns:
        pages = []
        for k, r in lines:
            pg = key_of(r[1].value)
            if not pages or pages[-1][0] != pg: pages.append((pg, []))
            pages[-1][1].append((k, r))
        bounds = []
        for pg, ls in pages:
            if pg in turns:
                cands = [p for p in (lists['coarse'] or lists['fine']) if abs(p[1] - turns[pg]) <= 3]
                bounds.append(min(cands, key=lambda p: abs(p[1] - turns[pg]))[1] if cands else turns[pg])
            else: bounds.append(None)
        for i in range(len(bounds)):
            if bounds[i] is None: bounds[i] = bounds[i - 1] if i else segs[0][0]
        for i, (pg, ls) in enumerate(pages):
            lo = bounds[i]; hi = bounds[i + 1] if i + 1 < len(pages) else dur
            gsegs = [(max(a_, lo), min(b_, hi)) for a_, b_ in segs if b_ > lo + 0.05 and a_ < hi - 0.05]
            groups.append((ls, gsegs))
    else:
        groups.append((lines, segs))

    yellow = PatternFill('solid', fgColor='FFF4B8')
    n = 0
    for ls, gsegs in groups:
        weights = [syllables(r[4].value, r[5].value) for k, r in ls]
        for (k, r), (s, e) in zip(ls, split_group(gsegs, weights)):
            if a.only_empty and r[3].value: continue
            if s is None: continue
            r[2].value, r[3].value = round(s, 2), round(e, 2); n += 1
            if len(r) > 7:
                note = (r[7].value or '')
                if 'Timing guessed' not in note: note = (note + ' Timing guessed; tap along in Time it mode.').strip()
                r[7].value = note; r[7].fill = yellow
    wb.save(xlsx)
    print(f'guessed times for {n} lines from {len(segs)} speech segments ({len(groups)} group{"s" if len(groups) > 1 else ""}); saved {os.path.basename(xlsx)}')


if __name__ == '__main__':
    main()
