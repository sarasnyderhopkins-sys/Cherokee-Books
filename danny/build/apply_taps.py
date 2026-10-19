#!/usr/bin/env python3
"""Put her tapped timings (2026-10-16) into the Sentences sheet.

Each tap is moved back to the start of the speech it reacted to (the end of the pause just before it),
and each line ends where the pause before the next line begins. Lines 49-51 are handled by hand (see notes).
"""
import re
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

TAPS = """1,10.67,14.23
2,16.2,18.2
3,20.06,21.42
4,23.33,34.46
5,36.33,39.97
6,40.22,46.44
7,48.79,53.56
8,55.05,61.44
9,63.3,69.84
10,70.83,74.2
11,75.6,85.26
12,86.57,87.76
13,89.21,94.08
14,96.77,99.67
15,101.87,105.92
16,109.03,110.36
17,111.68,116.61
18,117.7,123.1
19,124.29,129.62
20,130.28,132.26
21,133.79,135.96
22,136.94,143.65
23,144.71,150.69
24,150.73,155.93
25,157.43,161.63
26,163.61,168.08
27,169.96,171.9
28,173.91,179.02
29,180.54,182.42
30,183.22,187.63
31,189.16,195.09
32,196.63,203.94
33,205.7,209.35
34,209.35,212.42
35,213.39,216.04
36,217.53,223.75
37,225.24,231.6
38,234.16,237.66
39,240.22,243.44
40,245.1,252.21
41,253.17,255.32
42,257.01,263.38
43,265.26,267.32
44,268.73,269.62
45,271.1,274.01
46,275.46,285.45
47,287.62,288.78
48,289.95,293.19
49,293.84,308.4
50,309.35,314.61
51,314.68,324.85
title,0.92,10.42"""


def load(path):
    body = open(path, encoding='utf-8').read().split('\n')
    out = []
    for t in body[1].split():
        m = re.match(r'([\d.]+)-([\d.]+)\(', t)
        if m: out.append((float(m.group(1)), float(m.group(2))))
    out.sort()
    merged = []                      # touching pauses count as one
    for a, b in out:
        if merged and a - merged[-1][1] < 0.1: merged[-1] = (merged[-1][0], b)   # a 0.1 s blip does not end a pause
        else: merged.append((a, b))
    return merged

FINE = load('text/silences-fine.txt')      # -32 dB: finds every little pause (used to understand her taps)
COARSE = load('text/silences.txt')         # -35 dB
STRICT = load('text/silences-strict.txt')  # -38 dB: a pause has to be really quiet, so its edges hug the words less tightly
MARGIN = 0.15                              # extra lead before the first sound and extra tail after the last

def onset(tap, window=1.3):
    """Start of the speech she reacted to. A tap comes a little after the first sound, so look back for the
    end of the last pause (= start of a burst of speech) within `window` s; if the tap sits inside a pause with
    no such burst, the line starts when that pause ends."""
    prev = [s for s in FINE if s[1] <= tap + 0.05 and tap - s[1] <= window]
    if prev: return prev[-1][1], f'{tap - prev[-1][1]:.2f} s late'
    for s in FINE:
        if s[0] <= tap <= s[1]: return s[1], 'in pause'
    return round(tap - 0.4, 2), 'no pause nearby: tap minus 0.4 s reaction time'

def pause_at(t, lst):
    """The pause in lst that ends at (or just before) t, i.e. the pause before the speech starting at t."""
    c = [s for s in lst if s[0] - 0.05 <= t and s[1] <= t + 0.03 and t - s[1] <= 0.8]
    return c[-1] if c else None

def boundary(on):
    """(end of the previous line, start of this line) around a speech onset found in the fine list.
    Uses the quietest pause list that hears a pause in that gap (-38, then -35, then -32 dB), then widens by MARGIN on both sides."""
    pf = pause_at(on, FINE)
    pz = None
    for lst in (STRICT, COARSE):                        # the quietest list that still hears a pause in this gap
        q = pause_at(on, lst)
        if q and (pf is None or q[0] >= pf[0] - 0.05): pz = q; break
    pz = pz or pf
    if pz is None: return round(on - 0.25, 2), round(on - MARGIN, 2)
    a, b = pz
    m = min(MARGIN, max(0.03, (b - a) / 2 - 0.03))
    return round(a + m, 2), round(b - m, 2)

taps = {}
for row in TAPS.splitlines():
    k, a, b = row.split(',')
    taps[k] = (float(a), float(b))

onsets = {}
for k in map(str, range(1, 49)):
    t, how = onset(taps[k][0])
    onsets[k] = round(t, 2)
    print(f'{k:>3} tap {taps[k][0]:7.2f} -> {t:7.2f}  {how}')
# last page: her tap for 49 (293.84) is the "no usqualv..." clause inside line 48; the page turns at ~301.2 in the movie
onsets['49'] = 301.64
onsets['50'] = round(onset(taps['50'][0])[0], 2)
onsets['51'] = round(onset(taps['51'][0])[0], 2)

keys = [str(i) for i in range(1, 52)]
starts, ends = {}, {}
for i, k in enumerate(keys):
    e_prev, st = boundary(onsets[k])
    starts[k] = st
    if i > 0: ends[keys[i - 1]] = e_prev
# the last line ends where the closing pause begins, plus the margin
last = next(s for s in STRICT if s[0] > onsets['51'] + 1 and s[1] - s[0] > 5)
ends['51'] = round(last[0] + MARGIN, 2)
title_end = round(STRICT[0][0] + MARGIN, 2) if STRICT[0][0] < 6 else 4.5
print('title end', title_end, '| 51 ends', ends['51'])
for k in keys:
    if ends[k] <= starts[k]: raise SystemExit(f'line {k}: end {ends[k]} before start {starts[k]}')

LEAD = 0
wb = load_workbook('Trick or Treat Danny sentences.xlsx')
ws = wb['Sentences']
yellow = PatternFill('solid', fgColor='FFF4B8')
nofill = PatternFill(fill_type=None)
timing_note = re.compile(r'\s*(Timing guessed[^.]*\.|Page turn in the movie[^.]*\.(\s*start time may need a nudge\.)?|please check\.)', re.I)
for row in ws.iter_rows(min_row=2):
    k = str(row[0].value).strip().lower()
    k = str(int(float(k))) if re.fullmatch(r'\d+(\.0+)?', k) else k
    note = (row[7].value or '').strip()
    note = re.sub(r'\s*Page turn in the movie is [\d.]+ s from the nearest pause; start time may need a nudge\.', '', note)
    note = re.sub(r'\s*Timing guessed \(not enough pauses found\); please check\.', '', note)
    note = re.sub(r'\s*Timing: your tap for line 49.*?300\.75 s\.', '', note)
    note = re.sub(r'\s*Timing on this page is Claude’s best guess.*?last page\.', '', note)
    if k == 'title':
        row[2].value, row[3].value = 0, title_end
    elif k in starts:
        row[2].value = starts[k]
        row[3].value = ends[k]
        if k == '48':
            note = (note + ' Timing: your tap for line 49 (293.8 s) landed on “no usqualv…”, the second half of this sentence, so this line runs to the pause at 300.75 s.').strip()
        if k in ('49', '50', '51'):
            note = (note + ' Timing on this page is Claude’s best guess from the pauses (line 49 starts where the movie turns the page; 50 and 51 use your taps). Please listen to the last page.').strip()
    row[7].value = note or None
    row[7].fill = yellow if note else nofill
for row in wb['Settings'].iter_rows(min_row=2):
    if row[0].value == 'line_gap_seconds':
        row[1].value = 0
        row[2].value = 'Seconds taken off the end of the last line on a page (0: the ends already stop in the pause, so nothing is cut)'
wb.save('Trick or Treat Danny sentences.xlsx')
print()
for k in keys:
    print(k, starts[k], ends[k], f'({ends[k] - starts[k]:.2f} s)')
