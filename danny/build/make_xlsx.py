#!/usr/bin/env python3
"""Write 'Trick or Treat Danny sentences.xlsx' from sentences.json (first draft she edits)."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
data = json.load(open(os.path.join(HERE, 'sentences.json'), encoding='utf-8'))
sent = data['sentences']

bold = Font(bold=True)
yellow = PatternFill('solid', fgColor='FFF4B8')
wrap = Alignment(wrap_text=True, vertical='top')
chr_font = Font(name='Noto Sans Cherokee')

wb = Workbook()
ws = wb.active; ws.title = 'Sentences'
head = ['Line', 'Page', 'Start (s)', 'End (s)', 'Phonetics', 'Syllabary', 'English', 'Check notes']
ws.append(head)
ws.append(['title', 1, 0, 4.36, 'Didalonvhesgi nole Didane, Denili!', 'ᏗᏓᎶᏅᎮᏍᎩ ᏃᎴ ᏗᏓᏁ, ᏕᏂᎵ!', 'Trick or Treat, Danny!',
           'Syllabary typed from the cover art; the Word file has “Didalonvhesgi nole Didane Denili”. Plays over the cover when Begin is pressed (0–4.4 s of the narration).'])
for s in sent:
    ws.append([s['line'], s['page'], s['start'], s['end'], s['phon'], s['syl'], s['en'], s['note'].strip()])
for c in ws[1]: c.font = bold
widths = [6, 6, 9, 9, 46, 46, 46, 60]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = wrap
    row[5].font = chr_font
    if row[7].value: row[7].fill = yellow
ws.freeze_panes = 'A2'

# Pages
zones = {1: (0.72, 0.86), 5: (0.63, 0.89), 6: (0.73, 0.90), 7: (0.17, 0.25), 8: (0.76, 0.89), 9: (0.09, 0.31), 10: (0.13, 0.34), 11: (0.69, 0.91),
         12: (0.15, 0.28), 13: (0.23, 0.27), 14: (0.12, 0.21), 15: (0.12, 0.20), 16: (0.12, 0.30), 17: (0.72, 0.90), 18: (0.72, 0.90),
         19: (0.10, 0.27), 20: (0.72, 0.90), 21: (0.14, 0.31), 22: (0.11, 0.33), 23: (0.10, 0.23), 24: (0.74, 0.91)}
desc = {1: 'Front cover: Danny holding a jack-o’-lantern', 5: 'Danny at the window with a cold', 6: 'Dark windy night, moon behind a cloud',
        7: 'Danny whispers, sad', 8: 'Mother calls from the kitchen doorway', 9: 'Father, Danny and the pumpkin at the table',
        10: 'Danny scoops out the pumpkin', 11: 'Father lights the candle in the jack-o’-lantern', 12: 'Three monsters outside the window',
        13: 'Danny shouts for help', 14: 'The monsters come closer', 15: 'The ghost takes off its mask', 16: 'The friends jump up and down',
        17: 'Friends at the door, bags on the floor', 18: 'Father goes out with the bowl', 19: 'Friends take candy from the bowl',
        20: 'Danny in his mask at the window', 21: 'The friends pretend to be scared and wave', 22: 'Danny opens the bags on the table',
        23: 'Hot chocolate time', 24: 'The whole family hugs', 25: 'Back cover: the three friends in costume'}
wp = wb.create_sheet('Pages')
wp.append(['Page', 'Image file', 'Caption zone top', 'Caption zone bottom', 'Motion', 'Description'])
for p in [1] + list(range(5, 25)) + [25]:
    z = zones.get(p, ('', ''))
    wp.append([p, f'images/page-{p:02d}.jpg', z[0], z[1], 'hold' if p in (1, 25) else 'zoom in', desc.get(p, '')])
for c in wp[1]: c.font = bold
for i, w in enumerate([6, 20, 16, 18, 10, 50], 1): wp.column_dimensions[get_column_letter(i)].width = w

# Settings
wsett = wb.create_sheet('Settings')
wsett.append(['Setting', 'Value', 'What it does'])
rows = [
    ('title_syl', 'ᏗᏓᎶᏅᎮᏍᎩ ᏃᎴ ᏗᏓᏁ, ᏕᏂᎵ!', 'Title in the top bar (syllabary)'),
    ('title_phon', 'Didalonvhesgi nole Didane, Denili!', 'Title in the top bar (phonetics)'),
    ('title_en', 'Trick or Treat, Danny!', 'English title on the cover'),
    ('credit', 'Written by Edith Kunhardt · Translated by Wiggins “Inoli” Blackfox · Edited by Sara Snyder Hopkins · Translation © 2019', 'Credit line on the cover (from page 2 of the book). Add “Read by …” if you like.'),
    ('audio', 'audio/narration.mp3', 'The one narration file; Start and End on the Sentences sheet are seconds into it'),
    ('pause_seconds', 1.5, 'Pause after a page finishes before Play all turns the page'),
    ('text_default', 'on', 'on or off: is the text showing when the book opens'),
    ('show_syllabary', 'yes', 'yes/no: syllabary line in the caption'),
    ('show_phonetics', 'yes', 'yes/no: phonetics line in the caption'),
    ('show_english', 'yes', 'yes/no: English line in the caption'),
    ('fade_seconds', 0.6, 'How long the caption takes to fade in and out'),
    ('line_gap_seconds', 0.15, 'Taken off the end of each line so the next line’s first word is not clipped'),
]
for r in rows: wsett.append(list(r))
for c in wsett[1]: c.font = bold
for i, w in enumerate([18, 60, 70], 1): wsett.column_dimensions[get_column_letter(i)].width = w

# How to edit
wh = wb.create_sheet('How to edit')
for line in [
    'Trick or Treat, Danny! read-along: how to edit',
    '',
    'Sentences sheet: one row per sentence. Edit Phonetics, Syllabary and English freely. Page says which picture the line belongs to.',
    'Start and End are seconds into audio/narration.mp3. Lines on the same page play one after another from the first Start to the last End.',
    'Check notes (yellow) mark lines where the Oct 24 PDF and the Word file disagree, or where the timing was guessed. Clear a note once you have checked it.',
    '',
    'Easiest way to fix the timing: open index.html, press T to open timing mode, press Begin, then tap the space bar at the start of each line.',
    'When you reach the end it offers timings.csv. Save that file in this folder and rebuild; the page will use those times instead of the Start/End columns.',
    '',
    'Pages sheet: Caption zone top/bottom say where on the page (0 = top, 1 = bottom) the printed text used to be; the live caption sits there.',
    'Motion: zoom in, zoom out, hold. Description is the picture’s alt text.',
    '',
    'Rebuild after editing (Terminal, in this folder):',
    '    python3 build_book.py',
    '    python3 build_standalone.py        (optional: one file with everything inside)',
]:
    wh.append([line])
wh['A1'].font = bold
wh.column_dimensions['A'].width = 120

out = os.path.join(HERE, 'Trick or Treat Danny sentences.xlsx')
wb.save(out); print('wrote', out)
