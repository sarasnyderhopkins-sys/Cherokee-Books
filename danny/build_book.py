#!/usr/bin/env python3
"""Rebuild index.html from "Trick or Treat Danny sentences.xlsx" and template.html.

Run it from this folder after editing the spreadsheet:
    python3 build_book.py
If a timings.csv (saved by the page's Timing mode) sits in this folder, its Start/End times win.
"""
import csv, json, os, re, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'Trick or Treat Danny sentences.xlsx')
TEMPLATE = os.path.join(HERE, 'template.html')
OUT = os.path.join(HERE, 'index.html')
TIMINGS = os.path.join(HERE, 'timings.csv')


def cell(v):
    return '' if v is None else str(v).strip()


def num(v, default=None):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def main():
    wb = load_workbook(XLSX, data_only=True)

    settings = {}
    for row in wb['Settings'].iter_rows(min_row=2, values_only=True):
        if row and cell(row[0]):
            settings[cell(row[0])] = cell(row[1]) if len(row) > 1 else ''
    title = {'syl': settings.get('title_syl', ''), 'phon': settings.get('title_phon', ''), 'en': settings.get('title_en', '')}

    # Sentences: Line | Page | Start | End | Phonetics | Syllabary | English | notes
    lines, page_lines = {}, {}
    for row in wb['Sentences'].iter_rows(min_row=2, values_only=True):
        if not row or not cell(row[0]):
            continue
        key = cell(row[0]).lower()
        key = str(int(float(key))) if re.fullmatch(r'\d+(\.0+)?', key) else key
        page = cell(row[1])
        page = str(int(float(page))) if re.fullmatch(r'\d+(\.0+)?', page) else page
        lines[key] = {'syl': cell(row[5]), 'phon': cell(row[4]), 'en': cell(row[6]), 'start': num(row[2]), 'end': num(row[3]), 'page': page}
        page_lines.setdefault(page, []).append(key)

    if os.path.exists(TIMINGS):
        n = 0
        with open(TIMINGS, newline='', encoding='utf-8') as f:
            for r in csv.DictReader(f):
                k = cell(r.get('Line')); k = str(int(float(k))) if re.fullmatch(r'\d+(\.0+)?', k) else k
                if k in lines and num(r.get('Start')) is not None:
                    lines[k]['start'] = num(r['Start']); lines[k]['end'] = num(r['End'], lines[k]['end']); n += 1
        print(f'timings.csv: used tapped times for {n} lines')

    # Pages: Page | Image | zone top | zone bottom | Motion | Description
    pages = []
    for row in wb['Pages'].iter_rows(min_row=2, values_only=True):
        if not row or not cell(row[0]) or not re.fullmatch(r'\d+(\.0+)?', cell(row[0])):
            continue
        pid = str(int(float(cell(row[0]))))
        zt, zb = num(row[2]), num(row[3])
        pages.append({
            'id': pid,
            'image': cell(row[1]) or f'images/page-{int(pid):02d}.jpg',
            'zone': [zt, zb] if zt is not None and zb is not None else None,
            'motion': cell(row[4]) or 'hold',
            'alt': cell(row[5]),
            'lines': page_lines.get(pid, []),
        })
    pages.sort(key=lambda p: int(p['id']))

    # pauses in the narration, at three loudness thresholds (Timing mode puts starts and ends just outside them)
    def pause_list(*names):
        for cand in names:
            for folder in ('text', 'build/text'):
                path = os.path.join(HERE, folder, cand)
                if os.path.exists(path):
                    body = open(path, encoding='utf-8').read().split('\n')
                    out = []
                    for t in (body[1] if len(body) > 1 else '').split():
                        m = re.match(r'([\d.]+)-([\d.]+)\(', t)
                        if m: out.append([round(float(m.group(1)), 2), round(float(m.group(2)), 2)])
                    return out
        return []
    silences = pause_list('silences-fine.txt', 'silences.txt')          # -32 dB
    pauses = {'coarse': pause_list('silences.txt'), 'strict': pause_list('silences-strict.txt')}   # -35 dB, -38 dB

    missing = [p['image'] for p in pages if not os.path.exists(os.path.join(HERE, p['image']))]
    audio = settings.get('audio', 'audio/narration.mp3')
    if not os.path.exists(os.path.join(HERE, audio)): missing.append(audio)
    for m in missing: print(f'warning: file not found: {m}')
    untimed = [k for k, l in lines.items() if l['start'] is None or l['end'] is None]
    if untimed: print(f'warning: lines with no Start/End (they will be skipped): {untimed}')
    orphan = [p for p in page_lines if p not in {x['id'] for x in pages}]
    if orphan: print(f'warning: Sentences rows point at pages missing from the Pages sheet: {orphan}')

    data = {'settings': settings, 'title': title, 'pages': pages, 'lines': lines, 'silences': silences, 'pauses': pauses}
    html = open(TEMPLATE, encoding='utf-8').read()
    if '/*DATA*/' not in html: sys.exit('template.html has no /*DATA*/ marker')
    html = html.replace('/*DATA*/', json.dumps(data, ensure_ascii=False, indent=1).replace('</', '<\\/'))
    open(OUT, 'w', encoding='utf-8').write(html)
    print(f'wrote index.html: {len(pages)} pages, {len(lines)} lines')
    for p in pages:
        print(f'  page {p["id"]:>2}  {p["motion"]:<8} lines {", ".join(p["lines"]) or "-"}')


if __name__ == '__main__':
    main()
