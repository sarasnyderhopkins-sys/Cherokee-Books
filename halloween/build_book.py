#!/usr/bin/env python3
"""Rebuild index.html from "Halloween sentences.xlsx" and template.html.

Run it from this folder after editing the spreadsheet:
    python3 build_book.py
"""
import json, os, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'Halloween sentences.xlsx')
TEMPLATE = os.path.join(HERE, 'template.html')
OUT = os.path.join(HERE, 'index.html')


def cell(v):
    return '' if v is None else str(v).strip()


def main():
    wb = load_workbook(XLSX, data_only=True)

    # Settings sheet: key / value
    settings = {}
    for row in wb['Settings'].iter_rows(min_row=2, values_only=True):
        if row and cell(row[0]):
            settings[cell(row[0])] = cell(row[1]) if len(row) > 1 else ''

    # Sentences sheet: Line | Picture | Audio | Phonetics | Syllabary | English | notes
    title = {'syl': '', 'phon': '', 'en': '', 'subtitle': ''}
    lines = {}
    page_lines = {}
    for row in wb['Sentences'].iter_rows(min_row=2, values_only=True):
        if not row or not cell(row[0]):
            continue
        key = cell(row[0]).lower()
        pic = cell(row[1]).zfill(2) if cell(row[1]) else ''
        audio, phon, syl, en = (cell(x) for x in (row[2], row[3], row[4], row[5]))
        if key == 'title':
            title.update(syl=syl, phon=phon, en=en)
            continue
        if key == 'subtitle':
            title['subtitle'] = en or phon
            continue
        try:
            n = int(float(key))
        except ValueError:
            print(f'skipping row with Line = {key!r}')
            continue
        lines[str(n)] = {'syl': syl, 'phon': phon, 'en': en, 'audio': audio or f'audio/{n:02d}.mp3'}
        page_lines.setdefault(pic, []).append(str(n))

    # Pages sheet: Picture | Image | Lines | Motion | Effect | description
    pages = []
    for row in wb['Pages'].iter_rows(min_row=2, values_only=True):
        if not row or not cell(row[0]) or not cell(row[0]).isdigit():
            continue
        pid = cell(row[0]).zfill(2)
        pages.append({
            'id': pid,
            'image': cell(row[1]) or f'images/{pid}.jpg',
            'lines': page_lines.get(pid, []),
            'motion': cell(row[3]) or 'hold',
            'effect': cell(row[4]),
            'alt': cell(row[5]),
        })
    pages.sort(key=lambda p: p['id'])

    # Checks
    missing = [p['image'] for p in pages if not os.path.exists(os.path.join(HERE, p['image']))]
    missing += [l['audio'] for l in lines.values() if not os.path.exists(os.path.join(HERE, l['audio']))]
    for m in missing:
        print(f'warning: file not found: {m}')
    orphan = [n for n, ls in page_lines.items() if n and n not in {p['id'] for p in pages}]
    if orphan:
        print(f'warning: Sentences rows point at pictures that are not on the Pages sheet: {orphan}')

    data = {'settings': settings, 'title': title, 'pages': pages, 'lines': lines}
    html = open(TEMPLATE, encoding='utf-8').read()
    if '/*DATA*/' not in html:
        sys.exit('template.html has no /*DATA*/ marker')
    html = html.replace('/*DATA*/', json.dumps(data, ensure_ascii=False, indent=1).replace('</', '<\\/'))
    open(OUT, 'w', encoding='utf-8').write(html)
    print(f'wrote index.html: {len(pages)} pages, {len(lines)} lines')
    for p in pages:
        print(f'  {p["id"]}  {p["motion"]:<10} lines {", ".join(p["lines"]) or "-"}')


if __name__ == '__main__':
    main()
